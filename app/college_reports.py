"""تقارير الكلية: يجمعها وكيل الجودة تلقائياً من نماذج الأقسام (صف لكل قسم أو مقرر) بنفس أعمدة ملفات الإدارة،
ثم يصدّرها Excel ليرفعها على مستوى الكلية. لا يُدخل وكيل الجودة بيانات الأقسام مرة أخرى.
"""
from __future__ import annotations

from typing import Callable

from psycopg import AsyncConnection

from .form_templates import EXAMS, SEMESTERS, SKILLS_PERIODS, TEMPLATES
from .forms_engine import compute

STATUS_PCT = {"مُنفذ": 1.0, "مُنفذ جزئي": 0.5, "لم يُنفذ": 0.0}


def C(key, label, kind="text"):
    return {"key": key, "label": label, "type": kind}


def _rows_where(tables, section, field=None, value=None):
    rows = tables.get(section, [])
    return [r for r in rows if field is None or r.get(field) == value]


def _first(tables, section, field, value):
    xs = _rows_where(tables, section, field, value)
    return xs[0] if xs else None


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _share(rows, cols, eq):
    cells = [r.get(c) for r in rows for c in cols if r.get(c) not in (None, "")]
    return None if not cells else sum(1 for v in cells if v == eq) / len(cells)


# ------------------------------------------------------------------ بناة الصفوف: (القسم، الجداول المحسوبة، قيمة التصفية) → صفوف
def _exam_one(section_cols):
    def build(dept, t, f):
        r = _first(t, "exams", "exam", f)
        return [{"dept": dept, "semester": r.get("semester"), **{k: r.get(k) for k in section_cols}}] if r else []
    return build


def _analysis(dept, t, f):
    return [{"dept": dept, **r} for r in _rows_where(t, "courses", "exam", f)]


def _compare(dept, t, f):
    by = {}
    for r in t.get("courses", []):
        if r.get("exam") not in ("الفصلي الأول", "الفصلي الثاني") or not r.get("course"):
            continue
        g = by.setdefault((r.get("specialty"), r["course"]), {"dept": dept, "specialty": r.get("specialty"), "course": r["course"]})
        n = 1 if r["exam"] == "الفصلي الأول" else 2
        g.setdefault("max", r.get("max"))
        for k in ("registered", "high", "low", "below60", "absent", "deprived"):
            g[f"{k}{n}"] = r.get(k)
        for k in ("below60", "absent", "deprived"):
            g[f"{k}_pct{n}"] = r.get(f"{k}_pct")
    return list(by.values())


def _improve(section, field=None):
    def build(dept, t, f):
        return [{"dept": dept, **r} for r in t.get(section, []) if field is None or not f or r.get(field) in (f, None, "")]
    return build


def _by_semester(section, cols):
    def build(dept, t, f):
        r = _first(t, section, "semester", f)
        return [{"dept": dept, **{k: r.get(k) for k in cols}}] if r else []
    return build


def _outputs_exec(dept, t, f):
    n = 1 if f == "الأول" else 2
    rows = [r for r in (t.get("traits_test") or t.get("specialty_test") or []) if _num(r.get(f"target{n}")) is not None]
    if not rows:
        return []
    tgt = sum(_num(r.get(f"target{n}")) or 0 for r in rows)
    pres = sum(_num(r.get(f"present{n}")) or 0 for r in rows)
    return [{"dept": dept, "specialties": len(rows), "target": tgt, "date": None, "present": pres,
             "attend_pct": pres / tgt if tgt else None, "supervisor": None, "notes": None}]


def _plans_follow(dept, t, f):
    recs = [r for r in t.get("improvements", []) if r.get("recommendation")]
    if not recs:
        return []
    row = {"dept": dept}
    vals = []
    for i, r in enumerate(recs[:6], 1):
        v = STATUS_PCT.get(r.get("status"))
        row[f"r{i}"] = v
        if v is not None:
            vals.append(v)
    row["avg"] = sum(vals) / len(vals) if vals else None
    return [row]


def _course_desc(dept, t, f):
    n = 1 if f == "الأول" else 2
    trainers = [r for r in t.get("trainers", []) if r.get(f"file{n}") or r.get(f"lms{n}")]
    council = _first(t, "council", "semester", f)
    if not trainers and not council:
        return []
    return [{"dept": dept, "approved": council and council.get("approved"),
             "applied": _share(trainers, [f"file{n}"], "مكتمل"), "uploaded": _share(trainers, [f"lms{n}"], "مكتمل"),
             "trainers": len(trainers), "notes": council and council.get("meeting")}]


def _course_files(dept, t, f):
    return [{"dept": dept, **r} for r in _rows_where(t, "course_files", "semester", f)]


def _skills(dept, t, f):
    return [{"dept": dept, **r} for r in _rows_where(t, "sections", "period", f)]


def _tpl(code):
    return next(x for x in TEMPLATES if x["requirement_code"] == code)


_CF = _tpl("4-1-3")
_CF_COLS = next(s for s in _CF["schema"]["sections"] if s["key"] == "course_files")["columns"]

REPORTS: list[dict] = [
    # ---------------- الاختبارات
    {"key": "exam_readiness", "group": "الاختبارات", "source": "3-1/جاهزية", "filter": ("الاختبار", EXAMS),
     "title": "تقرير جاهزية المنشأة التدريبية لبدء الاختبار",
     "columns": [C("notify", "إبلاغ المتدربين بالأسبوع المحدد للاختبار"), C("courses", "عدد المقررات التدريبية", "num"),
                 C("aligned", "التأكد من أسئلة الاختبارات ومواءمتها مع الخطط"), C("materials", "كفاية العدد والخامات"),
                 C("halls", "جاهزية القاعات والورش والتجهيزات"), C("notes", "ملاحظات")],
     "build": _exam_one(["notify", "courses", "aligned", "materials", "halls", "notes"])},
    {"key": "exam_execution", "group": "الاختبارات", "source": "3-1/تنفيذ", "filter": ("الاختبار", EXAMS),
     "title": "تقرير تنفيذ الاختبار حسب اللوائح المنظمة",
     "columns": [C("ontime", "المقررات المختبرة في الوقت المحدد", "num"), C("moved", "المقررات المختبرة بتغيير الموعد", "num"),
                 C("theory", "المقررات النظرية", "num"), C("practical", "المقررات العملية", "num"),
                 C("absent", "عدد المتدربين الغائبين", "num"), C("recorded", "رصد النتائج"), C("feedback", "تغذية راجعة للمتدربين"),
                 C("notes", "ملاحظات")],
     "build": _exam_one(["ontime", "moved", "theory", "practical", "absent", "recorded", "feedback", "notes"])},
    {"key": "exam_analysis", "group": "الاختبارات", "source": "2-2-2", "filter": ("الاختبار", EXAMS),
     "title": "تقرير تحليل نتائج المتدربين للاختبار",
     "columns": [C("specialty", "التخصص"), C("course", "اسم المقرر"), C("hours", "عدد الساعات", "num"),
                 C("registered", "المسجلون حتى تاريخ الاختبار", "num"), C("max", "الدرجة القصوى", "num"),
                 C("high", "أعلى درجة", "num"), C("low", "أقل درجة", "num"),
                 C("below60", "أقل من 60٪ (عدد)", "num"), C("below60_pct", "أقل من 60٪ (نسبة)", "pct"),
                 C("absent", "الغائبون (عدد)", "num"), C("absent_pct", "الغائبون (نسبة)", "pct"),
                 C("deprived", "المحرومون (عدد)", "num"), C("deprived_pct", "المحرومون (نسبة)", "pct"),
                 C("method", "آلية القياس"), C("obstacles", "المعوقات والتحديات"), C("notes", "ملاحظات")],
     "build": _analysis},
    {"key": "exam_compare", "group": "الاختبارات", "source": "2-2-2", "filter": None,
     "title": "تقرير تحليل نتائج الاختبارات الفصلية والمقارنات",
     "columns": [C("specialty", "التخصص"), C("course", "المقرر")]
                + [C(f"{k}{n}", f"{lbl} — {ex}", kind) for k, lbl, kind in (
                    ("registered", "المسجلون", "num"), ("below60_pct", "أقل من 60٪", "pct"),
                    ("absent_pct", "الغياب", "pct"), ("deprived_pct", "الحرمان", "pct"))
                   for n, ex in ((1, "الفصلي الأول"), (2, "الفصلي الثاني"))],
     "build": _compare},
    {"key": "exam_improvements", "group": "الاختبارات", "source": "2-2-2", "filter": ("الاختبار", EXAMS),
     "title": "تقرير الإجراءات التحسينية لتحليل نتائج الاختبار",
     "columns": [C("recommendation", "توصية التحسين"), C("action", "الإجراء التصحيحي"), C("date", "تاريخ التنفيذ"),
                 C("owner", "مسؤول التنفيذ"), C("evidence", "شاهد التحسين"), C("status", "الإنجاز"), C("obstacles", "تحديات التنفيذ")],
     "build": _improve("improvements", "exam")},
    # ---------------- جودة المخرجات
    {"key": "outputs_requirements", "group": "قياس جودة المخرجات", "source": "2-2/المخرجات", "filter": ("الفصل", SEMESTERS),
     "title": "تقرير متطلبات قياس جودة مخرجات الكليات التقنية وفروعها",
     "columns": [C("specialties", "التخصصات المستهدفة", "num"), C("trainees", "المتدربون المستهدفون", "num"),
                 C("lab", "تهيئة المعمل"), C("internet", "تهيئة الإنترنت"), C("data_sent", "حصر وإرسال بيانات المتدربين"),
                 C("announced", "إعلان موعد الاختبار"), C("notes", "ملاحظات وتحديات")],
     "build": _by_semester("requirements", ["specialties", "trainees", "lab", "internet", "data_sent", "announced", "notes"])},
    {"key": "outputs_preparation", "group": "قياس جودة المخرجات", "source": "2-2/المخرجات", "filter": ("الفصل", SEMESTERS),
     "title": "تقرير تهيئة المتدربين لقياس جودة المخرجات",
     "columns": [C("specialties", "التخصصات المستهدفة", "num"), C("trainees", "المتدربون المستهدفون", "num"),
                 C("programs", "برامج التهيئة", "num"), C("attended", "الحضور", "num"), C("attend_pct", "نسبة الحضور", "pct"),
                 C("satisfaction", "قياس رضا المتدربين"), C("owner", "مسؤول التنفيذ"), C("notes", "ملاحظات وتحديات")],
     "build": _by_semester("preparation", ["specialties", "trainees", "programs", "attended", "attend_pct", "satisfaction", "owner", "notes"])},
    {"key": "outputs_execution", "group": "قياس جودة المخرجات", "source": "2-2-1", "filter": ("الفصل", SEMESTERS),
     "title": "تقرير تنفيذ قياس جودة مخرجات الكليات التقنية وفروعها",
     "columns": [C("specialties", "إجمالي التخصصات المستهدفة", "num"), C("target", "إجمالي المتدربين المستهدفين", "num"),
                 C("date", "تاريخ تنفيذ الاختبار"), C("present", "عدد الحضور من المتدربين", "num"),
                 C("attend_pct", "نسبة الحضور", "pct"), C("supervisor", "مشرف تنفيذ الاختبار"), C("notes", "ملاحظات وتحديات")],
     "build": _outputs_exec},
    {"key": "outputs_plans", "group": "قياس جودة المخرجات", "source": "2-2-1", "filter": None,
     "title": "تقرير متابعة تنفيذ خطط تحسين قياس جودة مخرجات التدريب",
     "columns": [C(f"r{i}", f"التوصية {o}", "pct") for i, o in
                 enumerate(["الأولى", "الثانية", "الثالثة", "الرابعة", "الخامسة", "السادسة"], 1)] + [C("avg", "النسبة العامة", "pct")],
     "build": _plans_follow},
    # ---------------- ملف المدرب وتوصيف المقرر
    {"key": "course_desc", "group": "ملف المدرب وتوصيف المقرر", "source": "4-1-3", "filter": ("الفصل", SEMESTERS),
     "title": "تقرير متابعة تطبيق المنشآت التدريبية لتوصيف المقرر التدريبي للمدرب",
     "columns": [C("approved", "اعتماد التوصيف من مجلس القسم"), C("applied", "تطبيق التوصيف من المدربين", "pct"),
                 C("uploaded", "رفع التوصيف على منصة التدريب الإلكتروني", "pct"), C("trainers", "عدد المدربين", "num"),
                 C("notes", "ملاحظات")],
     "build": _course_desc},
    {"key": "course_files", "group": "ملف المدرب وتوصيف المقرر", "source": "4-1-3", "filter": ("الفصل", SEMESTERS),
     "title": "نموذج متابعة ملف المقرر (نموذج رئيس القسم)",
     "columns": [C(c["key"], c["label"], "text") for c in _CF_COLS if c["key"] != "semester"],
     "build": _course_files},
    # ---------------- منصة المهارات العملية
    {"key": "skills", "group": "منصة المهارات العملية", "source": "4-1/المهارات", "filter": ("فترة التقرير", SKILLS_PERIODS),
     "title": "تقرير الإشراف ومتابعة التنفيذ – مشروع منصة المهارات",
     "columns": [C("specialty", "التخصص"), C("course", "المقرر التدريبي"), C("section", "رقم الشعبة"), C("trainer", "اسم المدرب"),
                 C("trainer_pct", "نسبة تفاعل المدربين مع أداة المهارات العملية (%)", "num"),
                 C("trainee_pct", "نسبة تفاعل المتدربين مع أداة المهارات العملية (%)", "num"), C("notes", "ملاحظات")],
     "build": _skills},
]
BY_KEY = {r["key"]: r for r in REPORTS}


# ------------------------------------------------------------------ تقارير من بيانات المنصة مباشرة (لا من النماذج)
SELF_AR = {"met": "متحقق", "partial": "متحقق جزئي", "not_met": "غير متحقق", "na": "لا ينطبق"}
WEEKS = ["الأول", "الثاني", "الثالث", "الرابع", "الخامس", "السادس", "السابع", "الثامن", "التاسع", "العاشر", "الحادي عشر",
         "الثاني عشر", "الثالث عشر", "الرابع عشر", "الخامس عشر", "السادس عشر", "السابع عشر", "الثامن عشر", "التاسع عشر"]


async def _patterns_rows(conn, cycle_id, flt, depts):
    cur = await conn.execute(
        """SELECT s.title_ar AS domain, ss.title_ar AS standard, r.code, r.text_ar, a.self_rating, a.ext_verdict,
                  (SELECT count(*) FROM assessment_evidence ae WHERE ae.assessment_id = a.id) AS ev
             FROM training_pattern_activation t JOIN assessment a ON a.cycle_id = t.patterns_cycle_id
             JOIN requirement r ON r.id = a.requirement_id JOIN sub_standard ss ON ss.id = r.sub_standard_id
             JOIN standard s ON s.id = ss.standard_id
            WHERE t.cycle_id = %s AND t.revoked_at IS NULL ORDER BY s.sort_order, ss.sort_order, r.sort_order""", (cycle_id,))
    return [{"dept": "الكلية", "domain": r["domain"], "standard": r["standard"], "requirement": f"{r['code']}  {r['text_ar']}",
             "level": SELF_AR.get(r["self_rating"]), "evidence": r["ev"]} for r in await cur.fetchall()]


async def _l2_rows(conn, cycle_id, flt, depts):
    """مستوى الجودة لكل معيار فرعي في كل قسم، ومستوى إنجاز إجراءاته التحسينية — يُحسب من التقديرات وخطط التحسين."""
    cur = await conn.execute(
        """SELECT d.name_ar AS dept, s.code AS std, s.title_ar AS std_title, ss.code AS sub, ss.title_ar AS sub_title,
                  array_agg(effective_rating(a.self_rating, a.ext_verdict, na_allowed(a.requirement_id, a.department_id))::text) AS ratings,
                  (SELECT count(*) FROM improvement_action ia JOIN assessment a2 ON a2.id = ia.assessment_id
                     JOIN requirement r2 ON r2.id = a2.requirement_id
                    WHERE a2.cycle_id = %(c)s AND a2.department_id = d.id AND r2.sub_standard_id = ss.id AND ia.status <> 'cancelled') AS plans,
                  (SELECT count(*) FROM improvement_action ia JOIN assessment a2 ON a2.id = ia.assessment_id
                     JOIN requirement r2 ON r2.id = a2.requirement_id
                    WHERE a2.cycle_id = %(c)s AND a2.department_id = d.id AND r2.sub_standard_id = ss.id AND ia.status = 'verified') AS closed
             FROM assessment a JOIN department d ON d.id = a.department_id
             JOIN requirement r ON r.id = a.requirement_id JOIN sub_standard ss ON ss.id = r.sub_standard_id
             JOIN standard s ON s.id = ss.standard_id
            WHERE a.cycle_id = %(c)s AND a.department_id = ANY(%(d)s) AND NOT is_excluded(a.requirement_id, a.department_id, a.cycle_id)
            GROUP BY d.id, d.name_ar, d.code, s.code, s.title_ar, s.sort_order, ss.id, ss.code, ss.title_ar, ss.sort_order
            ORDER BY d.code, s.sort_order, ss.sort_order""", {"c": cycle_id, "d": depts})
    out = []
    for r in await cur.fetchall():
        rs = [x for x in r["ratings"] if x]
        scored = [x for x in rs if x != "na"]
        if not rs:
            level = None
        elif not scored:
            level = "لا ينطبق"
        elif all(x == "met" for x in scored):
            level = "متحقق"
        elif all(x == "not_met" for x in scored):
            level = "غير متحقق"
        else:
            level = "متحقق جزئي"
        out.append({"dept": r["dept"], "std": f"{r['std']} {r['std_title']}", "std_code": r["std"], "sub": f"{r['sub']} {r['sub_title']}",
                    "level": level, "rated": f"{len(rs)} من {len(r['ratings'])}",
                    "actions": (r["closed"] / r["plans"]) if r["plans"] else None})
    return out


async def _visits(conn, cycle_id, flt, depts):
    cur = await conn.execute("SELECT * FROM external_visit WHERE cycle_id = %s ORDER BY visit_date, id", (cycle_id,))
    days = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
    return [{"dept": "الكلية", "id": v["id"], "date": v["visit_date"].isoformat(), "day": days[v["visit_date"].weekday()],
             "program": v["program"], "supervisor_name": v["supervisor_name"], "supervisor_phone": v["supervisor_phone"],
             "college_name": v["college_name"], "college_phone": v["college_phone"],
             "executed": None if v["executed"] is None else ("نعم" if v["executed"] else "لا"), "notes": v["notes_ar"]}
            for v in await cur.fetchall()]


REPORTS += [
    {"key": "l2_followup", "group": "نظام جودة التدريب (المستوى الثاني)", "source": None, "filter": ("الأسبوع", WEEKS),
     "title": "تقرير متابعة تنفيذ تحقيق نظام جودة التدريب في المنشآت التدريبية (المستوى الثاني)",
     "source_title": "تقديرات الأقسام وخطط التحسين في المنصة",
     "columns": [C("std", "المعيار الرئيسي"), C("sub", "المعيار الفرعي"), C("level", "مستوى الجودة"),
                 C("rated", "المتطلبات المقدّرة"), C("actions", "الإجراءات التحسينية (مستوى الإنجاز)", "pct")],
     "custom": _l2_rows},
    {"key": "external_visit", "group": "نظام جودة التدريب (المستوى الثاني)", "source": None, "filter": None,
     "title": "جدول زيارات التقييم الخارجي لنظام جودة التدريب", "source_title": "يُدخله وكيل الجودة",
     "columns": [C("date", "التاريخ"), C("day", "اليوم"), C("program", "البرنامج التدريبي"),
                 C("supervisor_name", "مشرف الجودة المكلف بالزيارة"), C("supervisor_phone", "جوال المشرف"),
                 C("college_name", "مشرف/وكيل الجودة في الكلية"), C("college_phone", "جوال وكيل الجودة"),
                 C("executed", "تنفيذ الزيارة"), C("notes", "ملاحظات")],
     "custom": _visits, "editable": True},
    {"key": "patterns_followup", "group": "أنماط التدريب", "source": None, "filter": ("الأسبوع", ["الثالث", "الخامس", "السابع", "التاسع"]),
     "title": "تقرير متابعة تنفيذ تحقيق منهجية قياس أنماط التدريب المختلفة", "source_title": "دورة منهجية قياس أنماط التدريب",
     "columns": [C("domain", "اسم المجال"), C("standard", "اسم المعيار"), C("requirement", "المتطلب"), C("level", "مستوى الجودة"),
                 C("evidence", "الشواهد المرفقة", "num")],
     "custom": _patterns_rows, "needs_patterns": True},
]
BY_KEY = {r["key"]: r for r in REPORTS}


def catalog() -> list[dict]:
    return [{"key": r["key"], "group": r["group"], "title": r["title"], "source": r["source"], "editable": r.get("editable", False),
             "needs_patterns": r.get("needs_patterns", False),
             "filter": r["filter"] and {"label": r["filter"][0], "options": r["filter"][1]}} for r in REPORTS]


async def build_report(conn: AsyncConnection, cycle_id: int, key: str, flt: str | None, dept_ids: list[int]) -> dict:
    rep = BY_KEY[key]
    if rep["filter"] and flt not in rep["filter"][1]:
        flt = rep["filter"][1][0]
    if rep.get("custom"):
        rows = await rep["custom"](conn, cycle_id, flt, dept_ids)
        return {"key": key, "title": rep["title"], "group": rep["group"], "filter": flt,
                "filter_label": rep["filter"] and rep["filter"][0], "columns": [C("dept", "القسم التدريبي")] + rep["columns"],
                "rows": rows, "total": None, "departments": [], "source": None, "source_title": rep["source_title"],
                "editable": rep.get("editable", False)}
    tpl = _tpl(rep["source"])
    cur = await conn.execute(
        """SELECT d.id, d.name_ar, f.id AS form_id, f.status, f.plan_data, f.follow_data
             FROM cycle_department cd JOIN department d ON d.id = cd.department_id
             LEFT JOIN (form_record f JOIN form_template t ON t.id = f.template_id AND t.requirement_code = %s)
                    ON f.cycle_id = cd.cycle_id AND f.department_id = d.id
            WHERE cd.cycle_id = %s AND d.id = ANY(%s) ORDER BY d.code""", (rep["source"], cycle_id, dept_ids))
    rows, depts = [], []
    for d in await cur.fetchall():
        got: list = []
        if d["form_id"]:
            tables = compute(tpl["schema"], d["plan_data"] or {}, d["follow_data"] or {})["tables"]
            got = rep["build"](d["name_ar"], tables, flt)
        rows.extend(got)
        depts.append({"id": d["id"], "name": d["name_ar"], "form_id": d["form_id"], "status": d["status"], "rows": len(got)})
    cols = [C("dept", "القسم التدريبي")] + rep["columns"]
    total = None
    if len(rows) > 1 and any(c["type"] in ("num", "pct") for c in rep["columns"]):
        total = {"dept": "إجمالي الكلية"}
        for c in rep["columns"]:
            vals = [_num(r.get(c["key"])) for r in rows if _num(r.get(c["key"])) is not None]
            if not vals or c["type"] not in ("num", "pct"):
                continue
            avg = c["type"] == "pct" or "%" in c["label"]          # النسب تُتوسط، والأعداد تُجمع
            total[c["key"]] = sum(vals) / len(vals) if avg else sum(vals)
        if "registered" in total:                                 # نسب الكلية من المجاميع (لا متوسط النسب)
            for k in ("below60", "absent", "deprived"):
                if total.get(k) is not None and total["registered"]:
                    total[f"{k}_pct"] = total[k] / total["registered"]
    return {"key": key, "title": rep["title"], "group": rep["group"], "filter": flt,
            "filter_label": rep["filter"] and rep["filter"][0], "columns": cols, "rows": rows, "total": total,
            "departments": depts, "source": rep["source"], "source_title": tpl["title_ar"]}


def to_xlsx(data: dict, institution: str, year: str, out_path: str, region: str | None = None) -> None:
    """ملف بهوية المؤسسة؛ ويُعبّأ في ملف الإدارة الأصلي متى توفر نموذجه."""
    from . import official_xlsx
    if official_xlsx.available(data["key"]):
        official_xlsx.fill(data["key"], data, institution, region, year, out_path)
        return
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from .xlsx_brand import body_cell, brand, header_cell
    wb = Workbook()
    ws = wb.active
    ws.title = data["title"][:30]
    cols = data["columns"]
    info = [("المنطقة", region or ""), ("المنشأة التدريبية", institution), ("العام التدريبي", year)]
    if data["filter"]:
        info.append((data["filter_label"], data["filter"]))
    head = brand(ws, data["title"], info, len(cols) + 1)
    for j, c in enumerate([{"label": "#"}] + cols, 1):
        header_cell(ws.cell(head, j, c["label"]))
        ws.column_dimensions[get_column_letter(j)].width = 6 if j == 1 else (14 if c.get("type") in ("num", "pct") else 26)
    ws.row_dimensions[head].height = 46
    body = data["rows"] + ([data["total"]] if data["total"] else [])
    for i, r in enumerate(body, 1):
        is_total = r is data["total"]
        body_cell(ws.cell(head + i, 1, "" if is_total else i), is_total)
        for j, c in enumerate(cols, 2):
            cell = ws.cell(head + i, j, r.get(c["key"]))
            body_cell(cell, is_total)
            if c["type"] == "pct" and isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%"
    ws.freeze_panes = ws.cell(head + 1, 2)
    wb.save(out_path)
