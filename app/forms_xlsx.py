"""نسخة Excel من تقارير النماذج ذات الجداول والأرقام — بهوية ملف PDF نفسها (الشعار، اسم الكلية، ألوان المنصة).

- ورقة «الملخص»: بيانات الوثيقة، المتطلب، المؤشرات (مربوطة بأوراق الجداول بصيغ)، النصوص، الشواهد، التواقيع.
- ورقة لكل جدول: جدول Excel (تصفية وفرز ورسوم بيانية مباشرة)، والأعمدة المحسوبة والملخصات صيغٌ حية
  تتحدث إذا عدّل المنسق الأرقام، وقوائم منسدلة لأعمدة الاختيار، وتلوين حالة التنفيذ.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .forms_engine import DONE, NOT_DONE, PARTIAL, _num
from .minutes_pdf import gregorian, hijri
from .pdf_brand import COLLEGE, ORG, ROOT, SYSTEM

FONT = "Segoe UI"
TEAL_D, TEAL, INK, SOFT, LINE = "126C66", "25AFA5", "1F2D2C", "EEF6F4", "CFDEDB"
THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_HEAD = PatternFill("solid", fgColor=TEAL_D)
F_GRP = PatternFill("solid", fgColor=TEAL)
F_SOFT = PatternFill("solid", fgColor=SOFT)
F_CALC = PatternFill("solid", fgColor="F4F8F7")
STATUS_FILL = {DONE: "D7F0E4", PARTIAL: "FFF1CC", NOT_DONE: "FBDADA"}
PCT, NUM = "0.0%", "#,##0.##"
HEADER_ROWS = 6                     # ترويسة الهوية في أعلى كل ورقة


def has_numbers(schema: dict) -> bool:
    """التقرير «فيه جداول وأرقام»: جدول بعمود رقمي أو محسوب، أو جدول تجميعي."""
    return any(s["type"] == "pivot" or (s["type"] == "table" and any(c["type"] in ("number", "computed") for c in s["columns"]))
               for s in schema["sections"])


def _font(size=10, bold=False, color=INK) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color)


def _sheet_name(title: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:28] or "جدول"
    name, k = base, 2
    while name in used:
        name, k = f"{base[:25]} ({k})", k + 1
    used.add(name)
    return name


def _q(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


def _brand(ws, ncols: int, kind: str, title: str, ref: str, status: str) -> None:
    """الترويسة: الشعار وأسفله اسم الكلية (يمين) — نوع الوثيقة وعنوانها (وسط) — المرجع والحالة (يسار)."""
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ncols = max(ncols, 6)
    for r in range(1, HEADER_ROWS):
        ws.row_dimensions[r].height = 21
    p = ROOT / "tvtc-logo.png"
    if p.is_file():
        img = XLImage(str(p))
        img.width, img.height = int(img.width * 74 / img.height), 74
        ws.add_image(img, "A1")
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=2)
    c = ws.cell(5, 1, COLLEGE)
    c.font, c.alignment = _font(10, True, TEAL_D), Alignment(horizontal="right", vertical="center")
    mid_a, mid_b = 3, max(3, ncols - 2)
    ws.merge_cells(start_row=1, start_column=mid_a, end_row=1, end_column=mid_b)
    ws.cell(1, mid_a, kind).font = _font(11, True, TEAL)
    ws.merge_cells(start_row=2, start_column=mid_a, end_row=3, end_column=mid_b)
    ws.cell(2, mid_a, title).font = _font(14, True, TEAL_D)
    ws.merge_cells(start_row=4, start_column=mid_a, end_row=4, end_column=mid_b)
    ws.cell(4, mid_a, f"{ORG} — {SYSTEM}").font = _font(9, color="5B6B69")
    for r in (1, 2, 4):
        ws.cell(r, mid_a).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    side = mid_b + 1
    for r, (txt, bold) in enumerate([(f"المرجع: {ref}", True), (f"تاريخ الإصدار: {datetime.now():%Y/%m/%d}م", False),
                                     (status, True)], start=1):
        if side < ncols:
            ws.merge_cells(start_row=r, start_column=side, end_row=r, end_column=ncols)
        cell = ws.cell(r, side, txt)
        cell.font = _font(9, bold, "B45309" if r == 3 and "مسودة" in status else (TEAL_D if r == 3 else INK))
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(1, ncols + 1):                     # خط فاصل بلون الهوية أسفل الترويسة
        ws.cell(HEADER_ROWS, col).border = Border(top=Side(style="medium", color=TEAL_D))
    ws.oddFooter.right.text = f"{ORG} · {COLLEGE}"
    ws.oddFooter.center.text = "صفحة &P من &N"
    ws.oddFooter.left.text = ref
    for f in (ws.oddFooter.right, ws.oddFooter.center, ws.oddFooter.left):
        f.size, f.color = 8, "5B6B69"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.top, ws.page_margins.bottom, ws.page_margins.left, ws.page_margins.right = 0.4, 0.5, 0.3, 0.3


def _head(cell, fill=F_HEAD) -> None:
    cell.fill, cell.border = fill, BOX
    cell.font = _font(10, True, "FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _body(cell, fmt: str | None = None, calc: bool = False, wrap: bool = True) -> None:
    cell.border = BOX
    cell.font = _font(10, calc)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    if calc:
        cell.fill = F_CALC
    if fmt:
        cell.number_format = fmt


def _section_title(ws, row: int, text: str, ncols: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(ncols, 2))
    c = ws.cell(row, 1, text)
    c.font, c.alignment = _font(12, True, TEAL_D), Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[row].height = 22


# ---------------------------------------------------------------------------
# الصيغ: تُترجم عمليات المحرك إلى صيغ Excel؛ وما لا مقابل له يُكتب قيمته المحسوبة
# ---------------------------------------------------------------------------
def _row_formula(expr: dict, ref) -> str | None:
    op = expr["op"]
    if op == "ratio":
        a, b = ref(expr["a"]), ref(expr["b"])
        return f'=IF(OR({a}="",N({b})=0),"",{a}/{b})' if a and b else None
    if op == "sub":
        a, b = ref(expr["a"]), ref(expr["b"])
        return f'=IF(OR({a}="",{b}=""),"",{a}-{b})' if a and b else None
    refs = [ref(x) for x in expr.get("of", [])] if isinstance(expr.get("of"), list) else []
    if refs and all(refs):
        args = ",".join(refs)
        if op == "add":
            return f'=IF(COUNT({args})=0,"",SUM({args}))'
        if op == "avg":
            return f'=IF(COUNT({args})=0,"",AVERAGE({args}))'
        if op == "mul" and not expr.get("map"):
            return f'=IF(COUNT({args})<{len(refs)},"",PRODUCT({args}))'
    if op == "band" and ref(expr["of"]):
        x = ref(expr["of"])
        v = f"TRUNC({x})" if expr.get("trunc") else x
        out = '""'
        for lo, hi, label in reversed(expr["bands"]):
            cond = f"AND({v}>={lo},{v}<{hi})" if hi is not None else f"{v}>={lo}"
            out = f"IF({cond},{_str(label)},{out})"
        return f'=IF({x}="","",{out})'
    if op == "map" and ref(expr["of"]) and all(len(str(a)) < 250 and len(str(b)) < 250 for a, b in expr["map"].items()):
        x, out = ref(expr["of"]), '""'
        for k, val in reversed(list(expr["map"].items())):
            out = f"IF({x}={_str(str(k))},{val if isinstance(val, (int, float)) else _str(str(val))},{out})"
        return "=" + out
    if op == "sumratio":
        nums, dens = [ref(x) for x in expr["num"]], [ref(x) for x in expr["den"]]
        if all(nums) and all(dens):
            n, d = ",".join(nums), ",".join(dens)
            return f'=IF(OR(COUNT({n})=0,SUM({d})=0),"",SUM({n})/SUM({d}))'
    return None


def _str(v: str) -> str:
    return '"' + v.replace('"', '""') + '"'


def _filled(rng: str) -> str:
    return f'SUMPRODUCT(--(LEN({rng})>0))'


def _summary_formula(expr: dict, rng, nrows: int) -> str | None:
    if nrows == 0:
        return None
    op = expr["op"]
    if op == "sum" and rng(expr["col"]):
        r = rng(expr["col"])
        return f'=IF(COUNT({r})=0,"",SUM({r}))'
    if op == "avg" and rng(expr["col"]):
        return f'=IFERROR(AVERAGE({rng(expr["col"])}),"")'
    if op == "count" and rng(None):
        return f"=ROWS({rng(None)})"
    if op == "count_eq" and rng(expr["col"]):
        return f'=COUNTIF({rng(expr["col"])},{_str(expr["eq"])})'
    if op == "ratio_sums" and not expr.get("where") and rng(expr["num"]) and rng(expr["den"]):
        n, d = rng(expr["num"]), rng(expr["den"])
        return f'=IF(OR(COUNT({n})=0,SUM({d})=0),"",SUM({n})/SUM({d}))'
    if op == "share_nonempty" and rng(expr["col"]):
        r = rng(expr["col"])
        return f'=IFERROR({_filled(r)}/ROWS({r}),"")'
    if op == "share_eq" and len(expr["cols"]) == 1 and rng(expr["cols"][0]):
        r = rng(expr["cols"][0])
        den = _filled(r) if expr.get("denominator") == "filled" else f"ROWS({r})"
        return f'=IFERROR(COUNTIF({r},{_str(expr["eq"])})/{den},"")'
    if op == "status_pct" and rng(expr["col"]):
        r = rng(expr["col"])
        done, part, nd = (f"COUNTIF({r},{_str(x)})" for x in (DONE, PARTIAL, NOT_DONE))
        den = f"ROWS({r})" if expr.get("denominator") == "rows" else f"({done}+{part}+{nd})"
        return f'=IFERROR(({done}+0.5*{part})/{den},"")'
    return None


def _value(c: dict, v):
    if v is None or v == "":
        return None
    if c["type"] == "number" or (c["type"] == "computed" and c.get("format") in ("pct", "num")):
        n = _num(v)
        return n if n is not None else v
    return v


def _fmt(c: dict) -> str | None:
    f = c.get("format") if c["type"] == "computed" else None
    if f == "pct":
        return PCT
    if c["type"] == "number" or f == "num":
        return NUM
    return None


# ---------------------------------------------------------------------------
def _table_sheet(wb, s: dict, rows: list[dict], sm: dict | None, ctx: dict, used: set[str], tno: int) -> dict:
    cols = s["columns"]
    ncols = len(cols) + 1
    ws = wb.create_sheet(_sheet_name(s.get("title", s["key"]), used))
    _brand(ws, ncols, ctx["kind"], s.get("title", s["key"]), ctx["ref"], ctx["status"])
    r = HEADER_ROWS + 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    info = ws.cell(r, 1, f"القسم التدريبي: {ctx['dept']}     ·     العام التدريبي: {ctx['year']}هـ     ·     المتطلب {ctx['code']}")
    info.font, info.fill, info.alignment = _font(10, True), F_SOFT, Alignment(horizontal="center", vertical="center")
    r += 2
    if any(c.get("group") for c in cols):                 # صف المجموعات (تخطيط / متابعة …)
        _head(ws.cell(r, 1), F_GRP)
        start, prev = 2, cols[0].get("group")
        for i, c in enumerate(cols + [{"group": object()}], start=2):
            g = c.get("group")
            if g != prev:
                if i - 1 > start:
                    ws.merge_cells(start_row=r, start_column=start, end_row=r, end_column=i - 1)
                cell = ws.cell(r, start, prev if isinstance(prev, str) else None)
                for k in range(start, i):
                    _head(ws.cell(r, k), F_GRP)
                start, prev = i, g
        r += 1
    hr = r
    labels, seen = ["م"], {"م"}
    for c in cols:
        lab, k = c["label"], 2
        while lab in seen:                                # رؤوس جدول Excel يجب أن تكون فريدة
            lab, k = f"{c['label']} ({k})", k + 1
        seen.add(lab)
        labels.append(lab)
    for j, lab in enumerate(labels, start=1):
        _head(ws.cell(hr, j, lab))
    ws.row_dimensions[hr].height = 34
    letter = {c["key"]: L(j) for j, c in enumerate(cols, start=2)}
    first, last = hr + 1, hr + max(len(rows), 1)
    for i, row in enumerate(rows):
        rr = first + i
        _body(ws.cell(rr, 1, i + 1))
        for j, c in enumerate(cols, start=2):
            v = None
            if c["type"] == "computed":
                v = _row_formula(c["expr"], lambda k: f"{letter[k]}{rr}" if k in letter else None)
            if v is None:
                v = _value(c, row.get(c["key"]))
            _body(ws.cell(rr, j, v), _fmt(c), calc=c["type"] == "computed")
    if not rows:
        _body(ws.cell(first, 1, 1))
        for j in range(2, ncols + 1):
            _body(ws.cell(first, j))
    # عرض الأعمدة
    ws.column_dimensions["A"].width = 6
    for j, c in enumerate(cols, start=2):
        longest = max([len(str(row.get(c["key"]) or "")) for row in rows] + [len(c["label"]) * 0.8, 8])
        ws.column_dimensions[L(j)].width = min(max(longest * 1.1, 10), 46)
    # جدول Excel: تصفية وفرز ونطاق جاهز للرسوم البيانية والحسابات
    tab = Table(displayName=f"T{tno}", ref=f"A{hr}:{L(ncols)}{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    # قوائم منسدلة وتلوين حالة التنفيذ لأعمدة الاختيار
    for j, c in enumerate(cols, start=2):
        opts = c.get("options") or []
        if c["type"] == "select" and opts and not any("," in o for o in opts) and len(",".join(opts)) < 250:
            dv = DataValidation(type="list", formula1=_str(",".join(opts)), allow_blank=True)
            dv.add(f"{L(j)}{first}:{L(j)}{last + 200}")
            ws.add_data_validation(dv)
        if c["type"] == "select" and DONE in opts:
            for val, color in STATUS_FILL.items():
                ws.conditional_formatting.add(f"{L(j)}{first}:{L(j)}{last + 200}",
                                              CellIsRule(operator="equal", formula=[_str(val)], fill=PatternFill("solid", fgColor=color)))
    ws.freeze_panes = ws.cell(hr + 1, 2)
    ws.print_title_rows = f"{hr}:{hr}"
    # الملخص أسفل الجدول — صيغ مرتبطة بالنطاق
    links: list[tuple[str, str, str | None]] = []
    if s.get("summary"):
        r = last + 2
        _section_title(ws, r, "ملخص الجدول", ncols)
        r += 1

        def rng(k):
            if k is None:
                return f"A{first}:A{last}"
            return f"{letter[k]}{first}:{letter[k]}{last}" if k in letter else None
        for m in s["summary"]:
            lc = ws.cell(r, 2, m["label"])
            if ncols >= 4:
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            _head(lc, F_GRP)
            vcol = 4 if ncols >= 4 else 3
            f = _summary_formula(m["expr"], rng, len(rows)) if m.get("expr") else None
            v = f if f else _value({"type": "computed", "format": m.get("format")}, (sm or {}).get(m["key"]))
            _body(ws.cell(r, vcol, v), PCT if m.get("format") == "pct" else NUM, calc=True)
            links.append((m["label"], f"={_q(ws.title)}!{L(vcol)}{r}", m.get("format")))
            r += 1
    return {"sheet": ws.title, "links": links}


def _pivot_sheet(wb, s: dict, groups: list[dict], ctx: dict, used: set[str], tno: int) -> None:
    cols = s["columns"]
    ncols = len(cols) + 1
    ws = wb.create_sheet(_sheet_name(s.get("title", s["key"]), used))
    _brand(ws, ncols, ctx["kind"], s.get("title", s["key"]), ctx["ref"], ctx["status"])
    hr = HEADER_ROWS + 2
    heads = [s.get("group_label", "التخصص"), "العدد"] + [c["label"] for c in cols]
    for j, lab in enumerate(heads, start=1):
        _head(ws.cell(hr, j, lab))
    ws.row_dimensions[hr].height = 34
    for i, g in enumerate(groups, start=hr + 1):
        _body(ws.cell(i, 1, str(g["group"])))
        _body(ws.cell(i, 2, g["count"]))
        for j, c in enumerate(cols, start=3):
            fmt = PCT if c.get("format") == "pct" else NUM
            _body(ws.cell(i, j, _value({"type": "computed", "format": c.get("format") or "num"}, g["values"].get(c["key"]))), fmt, calc=True)
    last = hr + max(len(groups), 1)
    if not groups:
        for j in range(1, len(heads) + 1):
            _body(ws.cell(hr + 1, j))
    ws.column_dimensions["A"].width = 30
    for j in range(2, len(heads) + 1):
        ws.column_dimensions[L(j)].width = 16
    tab = Table(displayName=f"T{tno}", ref=f"A{hr}:{L(len(heads))}{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = ws.cell(hr + 1, 2)


def render_xlsx(doc: dict) -> bytes:
    schema, calc = doc["schema"], doc["computed"]
    plan, follow = doc["plan_data"], doc["follow_data"]
    approved = doc["status"] == "approved"
    ctx = {"kind": "خطة" if doc.get("kind") == "plan" else "تقرير", "ref": f"AQAP-F-{doc['id']}",
           "status": "معتمد" if approved else "مسودة — غير معتمد", "dept": doc["department_name"],
           "year": doc["academic_year"], "code": doc["requirement_code"]}
    wb = Workbook()
    ws = wb.active
    ws.title = "الملخص"
    used = {ws.title}
    NC = 6
    _brand(ws, NC, ctx["kind"], doc["title_ar"], ctx["ref"], ctx["status"])
    for j, w in enumerate([22, 26, 22, 26, 22, 26], start=1):
        ws.column_dimensions[L(j)].width = w

    r = HEADER_ROWS + 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=NC)
    req = ws.cell(r, 1, f"المتطلب {doc['requirement_code']}: {doc.get('requirement_text') or ''}")
    req.font, req.fill = _font(10, True, TEAL_D), F_SOFT
    req.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[r].height = ws.row_dimensions[r + 1].height = 22
    r += 3
    meta = [("القسم التدريبي", doc["department_name"]), ("العام التدريبي", f"{doc['academic_year']}هـ"),
            ("نوع الوثيقة", ctx["kind"]), ("أعدّه", doc.get("submitted_by_name") or "—"),
            ("اعتمده", doc.get("approved_by_name") or "—"),
            ("تاريخ الاعتماد", gregorian(doc["approved_at"]) if doc.get("approved_at") else "—")]
    for i in range(0, len(meta), 3):
        for k, (lab, val) in enumerate(meta[i:i + 3]):
            lc, vc = ws.cell(r, 1 + 2 * k, lab), ws.cell(r, 2 + 2 * k, val)
            lc.fill, lc.font, lc.border, lc.alignment = F_SOFT, _font(10, True, TEAL_D), BOX, Alignment(horizontal="right", vertical="center")
            _body(vc)
        r += 1

    # أوراق الجداول، ثم مؤشراتها في الملخص بصيغ مرتبطة
    indicators: list[tuple[str, str, str | None, str]] = []
    tno = 1
    for s in schema["sections"]:
        k = s["key"]
        if s["type"] == "table":
            rows = calc["tables"].get(k, [])
            if k == "improvements" and not rows:
                continue
            res = _table_sheet(wb, s, rows, calc["summaries"].get(k), ctx, used, tno)
            indicators += [(lab, f, fmt, res["sheet"]) for lab, f, fmt in res["links"]]
            tno += 1
        elif s["type"] == "pivot":
            _pivot_sheet(wb, s, calc["pivots"].get(k, []), ctx, used, tno)
            tno += 1

    if indicators:
        r += 1
        _section_title(ws, r, "المؤشرات (تتحدث تلقائياً من أوراق الجداول)", NC)
        r += 1
        for j, lab in enumerate(["الجدول", "المؤشر", "القيمة"], start=1):
            _head(ws.cell(r, [1, 2, 4][j - 1], lab))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        _head(ws.cell(r, 3))
        r += 1
        for lab, f, fmt, sheet in indicators:
            _body(ws.cell(r, 1, sheet))
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
            _body(ws.cell(r, 2, lab))
            _body(ws.cell(r, 3))
            _body(ws.cell(r, 4, f), PCT if fmt == "pct" else NUM, calc=True)
            r += 1

    for s in schema["sections"]:                       # الحقول والنصوص
        k, title = s["key"], s.get("title", s["key"])
        if s["type"] == "fields":
            vals = {**(plan.get(k) or {}), **(follow.get(k) or {})}
            r += 1
            _section_title(ws, r, title, NC)
            r += 1
            for i, c in enumerate(s["fields"]):
                col = 1 + 2 * (i % 3)
                lc, vc = ws.cell(r, col, c["label"]), ws.cell(r, col + 1, _value(c, vals.get(c["key"])))
                lc.fill, lc.font, lc.border = F_SOFT, _font(10, True, TEAL_D), BOX
                lc.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
                _body(vc)
                if i % 3 == 2:
                    r += 1
            r += 1 if len(s["fields"]) % 3 else 0
        elif s["type"] == "text":
            v = (follow if s.get("phase") == "follow" else plan).get(k)
            r += 1
            _section_title(ws, r, title, NC)
            r += 1
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
            c = ws.cell(r, 1, v or "—")
            c.font, c.border = _font(10), BOX
            c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True, indent=1)
            ws.row_dimensions[r].height = max(20, 15 * (1 + len(v or "") // 110 + (v or "").count("\n")))
            r += 1

    atts = doc.get("attachments") or []
    if atts:
        r += 1
        _section_title(ws, r, "الشواهد والتقارير الداعمة", NC)
        r += 1
        for j, lab in enumerate(["م", "رمز الشاهد", "بيان الشاهد", "", "النوع", "الوصول"], start=1):
            _head(ws.cell(r, j, lab or None))
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        r += 1
        for i, a in enumerate(atts, 1):
            _body(ws.cell(r, 1, i))
            _body(ws.cell(r, 2, doc["requirement_code"]))
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
            _body(ws.cell(r, 3, (a.get("title_ar") or "") + (f" — {a['note']}" if a.get("note") else "")))
            _body(ws.cell(r, 4))
            _body(ws.cell(r, 5, "رابط" if a.get("kind") == "link" else "ملف"))
            acc = ws.cell(r, 6)
            if a.get("url"):
                acc.value, acc.hyperlink = "فتح الشاهد ↗", a["url"]
                _body(acc)
                acc.font = Font(name=FONT, size=10, color="0563C1", underline="single")
            else:
                _body(ws.cell(r, 6, f"بنك الشواهد: AQAP-E-{a.get('evidence_id') or a.get('id')}"))
            r += 1

    r += 1                                             # التواقيع
    if approved:
        boxes = [("أعدّه (منسق الجودة بالقسم)", doc.get("submitted_by_name") or ""),
                 ("اعتمده (رئيس القسم)", doc.get("approved_by_name") or ""),
                 ("تاريخ الاعتماد", f"{hijri(doc['approved_at'])} — {gregorian(doc['approved_at'])}")]
    else:
        boxes = [("أعدّه (منسق الجودة بالقسم)", ""), ("اعتمده (رئيس القسم)", ""), ("اطّلع عليه (وكالة الجودة)", "")]
    for k, (t, n) in enumerate(boxes):
        c1 = 1 + 2 * k
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c1 + 1)
        ws.merge_cells(start_row=r + 1, start_column=c1, end_row=r + 2, end_column=c1 + 1)
        h = ws.cell(r, c1, t)
        _head(h, F_GRP)
        ws.cell(r, c1 + 1).border = BOX
        b = ws.cell(r + 1, c1, (n or "الاسم: ……………………") + "\nالتوقيع: ……………………")
        b.font, b.alignment = _font(10), Alignment(horizontal="center", vertical="center", wrap_text=True)
        for rr in (r + 1, r + 2):
            for cc in (c1, c1 + 1):
                ws.cell(rr, cc).border = BOX
    ws.row_dimensions[r + 1].height = ws.row_dimensions[r + 2].height = 22
    if approved:
        r += 3
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
        e = ws.cell(r, 1, f"اعتُمد إلكترونياً عبر منصة الجودة والاعتماد — المرجع {ctx['ref']}؛ النسخة الإلكترونية هي المرجع.")
        e.font, e.alignment = _font(9, color=TEAL_D), Alignment(horizontal="center")

    wb.properties.title = doc["title_ar"]
    wb.properties.creator = f"{COLLEGE} — منصة الجودة والاعتماد"
    wb.properties.subject = f"{ctx['kind']} — المتطلب {doc['requirement_code']} — {doc['department_name']}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
