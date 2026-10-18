"""تصدير بديل بصيغة Excel منسّقة عندما لا يتوفر ملف القالب الرسمي (.xlsb) على الخادم.

أوراق الملف: ملخص الأقسام · الاستيفاء حسب المعيار · التقديرات التفصيلية · خطط التحسين.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from psycopg import AsyncConnection

SELF = {"met": "متحقق", "partial": "متحقق جزئي", "not_met": "غير متحقق", "na": "لا ينطبق"}
EXT = {"conforms": "منطبق", "not_conforming": "غير منطبق", "insufficient": "غير كافٍ"}
STATUS = {"draft": "مسودة", "submitted": "مرفوعة", "revision": "إعادة", "approved": "معتمدة",
          "in_progress": "قيد التنفيذ", "done": "منفذة", "verified": "مغلقة", "cancelled": "ملغاة"}
FILL = {"متحقق": "C6EFCE", "متحقق جزئي": "FFEB9C", "غير متحقق": "FFC7CE", "لا ينطبق": "E7E6E6"}
HEAD = PatternFill("solid", fgColor="0B5D46")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _sheet(wb, title, headers, widths, info=(), institution=""):
    from .xlsx_brand import brand, header_cell
    ws = wb.create_sheet(title)
    head = brand(ws, title, [x for x in info if x[0] != "المنشأة التدريبية"], len(headers),
                 next((v for k, v in info if k == "المنشأة التدريبية"), ""))
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        header_cell(ws.cell(head, i, h))
    ws.row_dimensions[head].height = 42
    ws.freeze_panes = ws.cell(head + 1, 1)
    return ws


def _rows(ws, rows):
    from .xlsx_brand import body_cell
    for r in rows:
        ws.append(r)
        for c in ws[ws.max_row]:
            body_cell(c)
            if isinstance(c.value, str) and c.value in FILL:
                c.fill = PatternFill("solid", fgColor=FILL[c.value])
            if isinstance(c.value, str) and len(c.value) > 40:
                c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)


async def build_basic(conn: AsyncConnection, cycle_id: int, out_path: str) -> dict:
    cur = await conn.execute("""SELECT c.title_ar, c.academic_year, i.name_ar AS inst, i.region_ar FROM assessment_cycle c
                                JOIN institution i ON i.id = c.institution_id WHERE c.id = %s""", (cycle_id,))
    cycle = await cur.fetchone()
    cur = await conn.execute(
        """SELECT COALESCE(d.name_ar, 'متطلبات المنشأة') AS dept, s.code AS sc, s.title_ar AS st, ss.code AS ssc,
                  ss.title_ar AS sst, r.code AS rc, r.text_ar AS rt, a.self_rating, a.ext_verdict, v.eff_rating, a.ext_recommendation,
                  is_excluded(a.requirement_id, a.department_id, a.cycle_id) AS excluded,
                  (SELECT count(*) FROM assessment_evidence ae WHERE ae.assessment_id = a.id) AS ev
             FROM assessment a JOIN requirement r ON r.id = a.requirement_id
             JOIN sub_standard ss ON ss.id = r.sub_standard_id JOIN standard s ON s.id = ss.standard_id
             LEFT JOIN department d ON d.id = a.department_id
             LEFT JOIN v_assessment_scored v ON v.assessment_id = a.id
            WHERE a.cycle_id = %s
            ORDER BY d.code NULLS FIRST, s.sort_order, ss.sort_order, r.sort_order""", (cycle_id,))
    rows = await cur.fetchall()
    cur = await conn.execute(
        """SELECT COALESCE(d.name_ar, 'متطلبات المنشأة') AS dept, t.requirement_code, t.requirement_text, t.action_ar,
                  COALESCE(t.owner_label, u.full_name_ar) AS owner, t.due_date, t.status, t.is_overdue
             FROM v_improvement_tracker t LEFT JOIN department d ON d.id = t.department_id
             LEFT JOIN app_user u ON u.id = t.owner_id
            WHERE t.cycle_id = %s AND t.status <> 'cancelled' ORDER BY d.code NULLS FIRST, t.requirement_code""", (cycle_id,))
    plans = await cur.fetchall()

    def score(xs):
        xs = [x for x in xs if not x["excluded"] and x["eff_rating"] in ("met", "partial", "not_met")]
        if not xs:
            return None
        return round(sum({"met": 1, "partial": .5, "not_met": 0}[x["eff_rating"]] for x in xs) / len(xs) * 100, 1)

    wb = Workbook()
    wb.remove(wb.active)
    info = [("المنطقة", (cycle or {}).get("region_ar") or ""), ("المنشأة التدريبية", (cycle or {}).get("inst") or ""),
            ("العام التدريبي", (cycle or {}).get("academic_year") or ""), ("الدورة", (cycle or {}).get("title_ar") or "")]
    depts = list(dict.fromkeys(r["dept"] for r in rows))
    ws = _sheet(wb, "ملخص الأقسام", ["القسم", "عدد المتطلبات", "متحقق", "متحقق جزئي", "غير متحقق", "لا ينطبق",
                                     "غير مقدّر", "بلا شواهد", "نسبة الاستيفاء %"], [26, 14, 11, 12, 12, 11, 11, 11, 16], info)
    out = []
    for d in depts:
        xs = [r for r in rows if r["dept"] == d and not r["excluded"]]
        c = lambda k: sum(1 for x in xs if x["eff_rating"] == k)
        out.append([d, len(xs), c("met"), c("partial"), c("not_met"), c("na"), c(None),
                    sum(1 for x in xs if not x["ev"] and x["eff_rating"] != "na"), score(xs)])
    _rows(ws, out)

    stds = list(dict.fromkeys((r["sc"], r["st"]) for r in rows))
    ws = _sheet(wb, "الاستيفاء حسب المعيار", ["القسم"] + [t for _, t in stds], [26] + [18] * len(stds), info)
    _rows(ws, [[d] + [score([r for r in rows if r["dept"] == d and r["sc"] == sc]) for sc, _ in stds] for d in depts])

    ws = _sheet(wb, "التقديرات التفصيلية", ["القسم", "المعيار", "المعيار الفرعي", "الرمز", "المتطلب", "التقييم الذاتي",
                                            "التقييم الخارجي", "التقدير المعتمد", "الشواهد", "توصية المقيّم"],
                [20, 20, 26, 9, 60, 13, 13, 13, 9, 40], info)
    _rows(ws, [[r["dept"], r["st"], r["sst"], r["rc"], r["rt"],
                "لا ينطبق" if r["excluded"] else SELF.get(r["self_rating"], ""), EXT.get(r["ext_verdict"], ""),
                "لا ينطبق" if r["excluded"] else SELF.get(r["eff_rating"], ""), r["ev"], r["ext_recommendation"] or ""]
               for r in rows])

    ws = _sheet(wb, "خطط التحسين", ["القسم", "الرمز", "المتطلب", "الإجراء", "المسؤول", "الموعد", "الحالة", "متأخرة"],
                [20, 9, 50, 50, 20, 12, 13, 9], info)
    _rows(ws, [[p["dept"], p["requirement_code"], p["requirement_text"], p["action_ar"], p["owner"] or "",
                p["due_date"].isoformat() if p["due_date"] else "", STATUS.get(p["status"], p["status"]),
                "نعم" if p["is_overdue"] else ""] for p in plans])
    wb.save(out_path)
    return {"mode": "basic", "departments": len(depts), "assessments": len(rows), "plans": len(plans),
            "cycle": cycle["title_ar"] if cycle else None}
