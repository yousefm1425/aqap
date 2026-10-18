"""هوية ملفات Excel المولّدة: شعار المؤسسة العامة للتدريب التقني والمهني، وشريط ترويسة بألوان نماذج الإدارة،
وتذييل «الإدارة العامة لجودة التدريب»، وإعدادات طباعة أفقية من اليمين لليسار."""
from __future__ import annotations

from pathlib import Path

from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LOGO = Path(__file__).parent / "report_templates" / "tvtc-logo.png"
BAND = PatternFill("solid", fgColor="DCE6F1")       # أزرق الترويسة في نماذج الإدارة
HEAD = PatternFill("solid", fgColor="BDD7EE")
LABEL = PatternFill("solid", fgColor="D9D9D9")
TOTAL = PatternFill("solid", fgColor="FFF2CC")
INK = "1F3864"
THIN = Side(style="thin", color="8EA9C8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Sakkal Majalla"


def brand(ws, title: str, info: list[tuple[str, str]], ncols: int) -> int:
    """يرسم الترويسة ويعيد رقم صف رؤوس الأعمدة."""
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ncols = max(ncols, 6)
    last = get_column_letter(ncols)
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 24
        for c in range(1, ncols + 1):
            ws.cell(r, c).fill = BAND
    if LOGO.is_file():
        img = XLImage(str(LOGO))
        img.height, img.width = 62, int(62 * 1620 / 396)
        ws.add_image(img, "A1")
    tcol = min(4, ncols - 2)
    ws.merge_cells(start_row=1, start_column=tcol, end_row=2, end_column=ncols)
    t = ws.cell(1, tcol, title)
    t.font = Font(name=FONT, bold=True, size=18, color=INK)
    t.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells(start_row=3, start_column=tcol, end_row=3, end_column=ncols)
    s = ws.cell(3, tcol, "الإدارة العامة لجودة التدريب")
    s.font = Font(name=FONT, size=12, color=INK)
    s.alignment = Alignment(horizontal="center")
    # بيانات الترويسة في سطر واحد: المنطقة · المنشأة · العام ...
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=ncols)
    line = ws.cell(5, 1, "     ·     ".join(f"{k}: {v}" for k, v in info if v not in (None, "")))
    line.font = Font(name=FONT, bold=True, size=12, color=INK)
    line.fill, line.alignment = LABEL, Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 22
    row = 5
    ws.oddFooter.right.text = "الإدارة العامة لجودة التدريب"
    ws.oddFooter.left.text = "صفحة &P من &N"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    _ = last
    return row + 2


def header_cell(cell) -> None:
    cell.fill, cell.border = HEAD, BOX
    cell.font = Font(name=FONT, bold=True, size=12, color=INK)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def body_cell(cell, total: bool = False) -> None:
    cell.border = BOX
    cell.font = Font(name=FONT, size=12, bold=total, color=INK if total else "000000")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if total:
        cell.fill = TOTAL
