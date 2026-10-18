"""تعبئة تقارير الكلية في ملفات الإدارة الأصلية: يبقى الهيدر وشعار المؤسسة والألوان والتنسيق كما هي،
وتُكتب البيانات في خلاياها. إن تجاوزت الصفوف سعة النموذج يُنسخ تنسيق آخر صف.
"""
from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

TPL_DIR = Path(__file__).parent / "report_templates"
STATUS_PCT = {"مُنفذ": 1.0, "مُنفذ جزئي": 0.5, "لم يُنفذ": 0.0}
AUDIT_TPL = {"قيّم - يعتمد": "قيم - يعتمد", "قيّم - يراجع": "قيم-يراجع", "لم يقيّم": "لم يقيم"}
L = "long"     # «الفصل التدريبي الأول»
S = "short"    # «الأول»

ITEMS = ["trainer_data", "contact", "course_data", "desc", "goal", "reqs", "plan", "refs", "quality", "links"]

# لكل تقرير: الورقة، خلايا الترويسة، صف البداية، عمود المنشأة، وأعمدة الحقول
MAP: dict[str, dict] = {
    "exam_readiness": {"region": "G9", "semester": ("G10", L), "filter": "L9", "year": "M10", "row": 15,
                       "cols": {"notify": "E", "courses": "G", "aligned": "I", "materials": "K", "halls": "M"}},
    "exam_execution": {"region": "G9", "semester": ("G10", L), "filter": "N9", "year": "O10", "row": 15,
                       "cols": {"ontime": "E", "moved": "G", "theory": "H", "practical": "J", "absent": "K",
                                "recorded": "M", "feedback": "O"}},
    "exam_analysis": {"region": "G9", "semester": ("G10", L), "filter": "S9", "year": "T10", "row": 15, "dept": "E",
                      "cols": {"course": "F", "hours": "G", "registered": "H", "max": "I", "high": "J", "low": "K",
                               "below60": "L", "absent": "N", "deprived": "P", "method": "R", "obstacles": "S", "notes": "T"}},
    "exam_compare": {"region": "G9", "semester": ("G10", L), "year": "AE10", "row": 16, "dept": "E",
                     "cols": {"course": "F", "registered1": "H", "registered2": "I", "max": "J", "high1": "K", "high2": "L",
                              "low1": "M", "low2": "N", "below601": "O", "below602": "Q", "absent1": "S", "absent2": "U",
                              "deprived1": "W", "deprived2": "Y"}},
    "exam_improvements": {"region": "G9", "semester": ("G10", L), "filter": "N9", "year": "O10", "row": 14,
                          "cols": {"recommendation": "E", "action": "H", "date": "K", "owner": "L", "evidence": "M",
                                   "progress": "N", "obstacles": "O"}},
    "outputs_requirements": {"region": "F9", "semester": ("F10", L), "year": "M10", "row": 15, "by_filter": True,
                             "cols": {"specialties": "E", "trainees": "F", "lab": "G", "internet": "H", "data_sent": "I",
                                      "announced": "K", "notes": "M"}},
    "outputs_preparation": {"region": "F9", "semester": ("F10", S), "year": "M10", "row": 15, "by_filter": True,
                            "cols": {"specialties": "E", "trainees": "F", "programs": "G", "attended": "H",
                                     "satisfaction": "I", "owner": "K", "notes": "M"}},
    "outputs_execution": {"region": "F9", "semester": ("F10", S), "year": "M10", "row": 15, "by_filter": True,
                          "cols": {"specialties": "E", "target": "F", "present": "I", "supervisor": "K", "notes": "M"}},
    "outputs_plans": {"region": "F9", "year": "K10", "row": 15, "count": "H10",
                      "cols": {"r1": "E", "r2": "F", "r3": "G", "r4": "H", "r5": "I", "r6": "J"}},
    "course_desc": {"region": "E9", "semester": ("E10", S), "year": "K10", "row": 15, "by_filter": True,
                    "cols": {"approved": "E", "applied": "G", "uploaded": "I", "notes": "K"}},
    "course_files": {"sheet": "نموذج رئيس القسم", "region": "Q1", "semester": ("L4", S), "year": "E4", "year_plain": True,
                     "row": 8, "dept": "D", "clear": "GHIJKLMNOPQRST", "by_filter": True,
                     "cols": {"trainer": "E", "course": "F", "link": "G", **dict(zip(ITEMS, "HIJKLMNOPQ")),
                              "audit_hod": "R", "audit_dean": "S", "audit_region": "T"}},
    "skills": {"region": "F7", "year": "F8", "filter": "F11", "row": 15,
               "cols": {"specialty": "D", "course": "E", "section": "F", "trainer": "G", "trainer_pct": "H",
                        "trainee_pct": "I", "notes": "J"}},
    "patterns_followup": {"region": "G9", "filter": "J9", "year": "K10", "row": 14, "fixed": True, "cols": {"level": "J"}},
    "l2_followup": {"region": "F9", "filter": "I9", "year": "J10", "row": 14, "dept": "D",
                    "cols": {"std_code": "E", "sub": "F", "level": "H", "actions": "J"}},
    "external_visit": {"region": "G9", "year": "M10", "row": 15,
                       "cols": {"date": "E", "day": "F", "bach": "G", "dipl": "H", "supervisor_name": "I", "supervisor_phone": "J",
                                "college_name": "K", "college_phone": "L", "executed": "M"}},
}


def available(key: str) -> bool:
    return key in MAP and (TPL_DIR / f"{key}.xlsx").is_file()


def _put(ws, ref: str, value) -> None:
    cell = ws[ref]
    if isinstance(cell, MergedCell):           # الكتابة في الخلية الأولى من النطاق المدمج
        for rng in ws.merged_cells.ranges:
            if ref in rng:
                cell = ws.cell(rng.min_row, rng.min_col)
                break
    if isinstance(cell.value, str) and cell.value.startswith("="):
        return                                 # لا تُستبدل صيغ النموذج
    cell.value = value


def _copy_row_style(ws, src: int, dst: int) -> None:
    for c in range(1, ws.max_column + 1):
        a, b = ws.cell(src, c), ws.cell(dst, c)
        if a.has_style:
            b._style = copy(a._style)
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def _value(key: str, field: str, r: dict):
    v = r.get(field)
    if key == "exam_improvements" and field == "progress":
        return STATUS_PCT.get(r.get("status"))
    if key == "course_desc" and field in ("applied", "uploaded") and isinstance(v, (int, float)):
        return "نعم" if v >= 0.999 else ("لا" if v == 0 else f"جزئي ({round(v * 100)}%)")
    if key == "course_files" and field.startswith("audit_"):
        return AUDIT_TPL.get(v, v)
    if key == "l2_followup" and field == "actions":
        return "لا ينطبق" if v is None else "متحقق" if v >= 0.999 else "غير متحقق" if v == 0 else "متحقق جزئي"
    if key == "external_visit" and field in ("bach", "dipl"):
        return "√" if (("بكالوريوس" if field == "bach" else "دبلوم") in (r.get("program") or "")) else None
    return v


def _college_under_logo(ws, key: str, institution: str) -> None:
    """يكتب اسم الكلية أسفل شعار المؤسسة العامة للتدريب التقني والمهني مباشرة."""
    from openpyxl.styles import Alignment, Font
    if not ws._images:
        return
    anc = ws._images[0].anchor
    c1, c2 = anc._from.col + 1, max(anc.to.col + 1, anc._from.col + 3)
    if key == "course_files":                    # الشعار في الصف الأول؛ الاسم في الصف الثاني
        row, c1, c2 = 2, 2, 5
    elif key == "skills":                        # رفع الشعار للصفوف 1-3 وكتابة الاسم في الصف الرابع
        for r in (1, 2, 3):
            ws.row_dimensions[r].height = 30
        anc._from.row, anc._from.rowOff = 0, 0
        anc.to.row, anc.to.rowOff = 3, 0
        row, c1, c2 = 4, 3, 9
    else:
        row = max(anc.to.row + 2, 6)
    for rng in list(ws.merged_cells.ranges):     # لا تداخل مع دمج قائم في الصف
        if rng.min_row <= row <= rng.max_row and not (rng.max_col < c1 or rng.min_col > c2):
            return
    ref = ws.cell(row, c1)
    above = ws.cell(max(row - 1, 1), c1)
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    ref.value = institution
    ref.font = Font(name="Sakkal Majalla", bold=True, size=14, color="1F3864")
    ref.alignment = Alignment(horizontal="center", vertical="center")
    if above.fill and above.fill.fill_type:
        ref.fill = copy(above.fill)
    ws.row_dimensions[row].height = 24


def fill(key: str, data: dict, institution: str, region: str | None, year: str, out_path: str) -> None:
    m = MAP[key]
    wb = load_workbook(TPL_DIR / f"{key}.xlsx")
    ws = wb[m["sheet"]] if m.get("sheet") else wb.worksheets[0]
    _college_under_logo(ws, key, institution)
    y = int(year) if str(year).isdigit() else None
    if region and m.get("region"):
        _put(ws, m["region"], region)
    if m.get("year"):
        _put(ws, m["year"], y if m.get("year_plain") else (f"{y}/{y + 1} هـ" if y else year))
    flt = data.get("filter")
    sem = flt if m.get("by_filter") else next((r.get("semester") for r in data["rows"] if r.get("semester")), None)
    if m.get("semester") and sem:
        ref, style = m["semester"]
        _put(ws, ref, f"الفصل التدريبي {sem}" if style == L else sem)
    if m.get("filter") and flt:
        _put(ws, m["filter"], flt)
    if m.get("count"):
        _put(ws, m["count"], 1)

    rows, start = data["rows"], m["row"]
    capacity = ws.max_row - start + 1
    for i, r in enumerate(rows):
        n = start + i
        if i >= capacity and not m.get("fixed"):
            _copy_row_style(ws, start + capacity - 1, n)
            _put(ws, f"B{n}", i + 1)
        if not m.get("fixed"):
            dept = r.get("dept")
            if m.get("dept"):
                _put(ws, f"C{n}", institution)
                _put(ws, f"{m['dept']}{n}", dept)
            elif m.get("dept", "C") is not None:
                _put(ws, f"C{n}", institution if dept in (None, "الكلية") else f"{institution} — {dept}")
        else:
            _put(ws, f"C{n}", institution)
        for field, col in m["cols"].items():
            v = _value(key, field, r)
            if v is not None and v != "":
                _put(ws, f"{col}{n}", v)
    if m.get("clear"):                         # إزالة عبارات «قيم هنا / رابط» من الصفوف غير المستخدمة
        for n in range(start + len(rows), ws.max_row + 1):
            for col in m["clear"]:
                c = ws[f"{col}{n}"]
                if not isinstance(c, MergedCell) and c.value in ("قيم هنا", "رابط"):
                    c.value = None
    wb.save(out_path)
