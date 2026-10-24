"""نسخة Excel للتقارير: تُولَّد لكل قالب فيه أرقام، بصيغ حية وجدول Excel لكل جدول."""
import datetime
import io

from openpyxl import load_workbook

from app.form_templates import TEMPLATES
from app.forms_engine import compute, empty_plan
from app.forms_xlsx import has_numbers, render_xlsx


def _doc(t: dict) -> dict:
    sch = t["schema"]
    plan = empty_plan(sch)
    for s in sch["sections"]:
        if s["type"] != "table":
            continue
        rows = plan.get(s["key"]) or [{} for _ in range(2)]
        out = []
        for i, r in enumerate(rows):
            r = {"id": f"r{i}", **r}
            for c in s["columns"]:
                if c["type"] == "number":
                    r.setdefault(c["key"], 5 + i)
                elif c["type"] == "select" and c.get("options"):
                    r.setdefault(c["key"], c["options"][0])
            out.append(r)
        plan[s["key"]] = out
    return dict(id=1, schema=sch, computed=compute(sch, plan, {}), plan_data=plan, follow_data={}, status="approved",
                kind=t.get("kind"), title_ar=t["title_ar"], requirement_code=t["requirement_code"], requirement_text="نص",
                department_name="تقنية الحاسب", academic_year=1447, submitted_by_name="أ", approved_by_name="ب",
                approved_at=datetime.datetime(2026, 10, 1), attachments=[{"title_ar": "شاهد", "kind": "link", "url": "https://example.com", "evidence_id": 1}])


def test_xlsx_for_every_numeric_template():
    numeric = [t for t in TEMPLATES if has_numbers(t["schema"])]
    assert numeric
    for t in numeric:
        doc = _doc(t)
        wb = load_workbook(io.BytesIO(render_xlsx(doc)))
        assert wb.sheetnames[0] == "الملخص"
        assert wb["الملخص"].sheet_view.rightToLeft
        tables = sum(len(ws.tables) for ws in wb)
        assert tables >= 1, t["requirement_code"]
        calc = any(s.get("summary") or any(c["type"] == "computed" for c in s.get("columns", [])) for s in t["schema"]["sections"]
                   if s["type"] == "table" and doc["computed"]["tables"].get(s["key"]))
        assert not calc or any(isinstance(c.value, str) and c.value.startswith("=") for ws in wb for row in ws.iter_rows() for c in row), t["requirement_code"]
