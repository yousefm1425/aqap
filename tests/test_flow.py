"""سيناريو كامل: فتح الدورة ← شاهد ← تقييم ذاتي ← اعتماد ← تدقيق ← نهائي ← تقارير ← خطة تحسين."""
API = "/api/v1"
COORD, HOD, AUDITOR, DEAN, COORD_EE = "coord.cs@ynb", "hod.cs@ynb", "auditor@ynb", "dean@ynb", "coord.ee@ynb"
STATE: dict = {}


def test_auth_required(client):
    assert client.get(f"{API}/me").status_code == 401
    assert client.get(f"{API}/me", headers={"Authorization": "Bearer junk"}).status_code == 401


def test_only_dean_opens_cycle(client, auth):
    assert client.post(f"{API}/cycles/1/open", headers=auth(COORD)).status_code == 403
    r = client.post(f"{API}/cycles/1/open", headers=auth(DEAN))
    assert r.status_code == 200, r.text
    assert r.json()["assessments_generated"] == 7          # 3 متطلبات × قسمين + متطلب منشأة


def test_department_scoping(client, auth):
    cs = client.get(f"{API}/cycles/1/assessments", headers=auth(COORD)).json()
    assert {a["department_id"] for a in cs} == {1}
    ee = client.get(f"{API}/cycles/1/assessments", headers=auth(COORD_EE)).json()
    assert {a["department_id"] for a in ee} == {2}
    assert len(client.get(f"{API}/cycles/1/assessments", headers=auth(DEAN)).json()) == 7
    STATE["a1"] = next(a["id"] for a in cs if a["requirement_code"] == "R1")
    STATE["ee_any"] = ee[0]["id"]
    assert client.get(f"{API}/assessments/{STATE['ee_any']}", headers=auth(COORD)).status_code == 404


def test_submit_requires_evidence(client, auth):
    a1 = STATE["a1"]
    r = client.patch(f"{API}/assessments/{a1}/self", headers=auth(COORD),
                     json={"self_rating": "met"})
    assert r.status_code == 200 and r.json()["self_by"] == 1
    r = client.post(f"{API}/assessments/{a1}/transition", headers=auth(COORD), json={"to_status": "submitted"})
    assert r.status_code == 409 and "شاهد" in r.json()["detail"]


def test_evidence_upload_and_link(client, auth):
    r = client.post(f"{API}/evidence", headers=auth(COORD),
                    json={"title_ar": "الهيكل التنظيمي للقسم", "kind": "file", "department_id": 1})
    assert r.status_code == 201, r.text
    ev = r.json()["id"]
    STATE["ev"] = ev
    bad = client.post(f"{API}/evidence/{ev}/versions", headers=auth(COORD), files={"file": ("x.exe", b"MZ")})
    assert bad.status_code == 415
    for content in (b"%PDF-1.4 v1", b"%PDF-1.4 v2"):
        r = client.post(f"{API}/evidence/{ev}/versions", headers=auth(COORD),
                        files={"file": ("الهيكل.pdf", content, "application/pdf")})
        assert r.status_code == 201, r.text
    assert r.json()["version_no"] == 2 and len(r.json()["sha256"]) == 64
    assert client.post(f"{API}/evidence/{ev}/versions", headers=auth(COORD_EE),
                       files={"file": ("a.pdf", b"x")}).status_code == 404
    r = client.post(f"{API}/assessments/{STATE['a1']}/evidence", headers=auth(COORD), json={"evidence_id": ev})
    assert r.status_code == 201, r.text
    d = client.get(f"{API}/evidence/{ev}/download", headers=auth(HOD)).json()
    assert d["url"].startswith("file://")


def test_full_workflow(client, auth):
    a1 = STATE["a1"]
    t = lambda who, to, **kw: client.post(f"{API}/assessments/{a1}/transition", headers=auth(who),
                                          json={"to_status": to, **kw})
    assert t(COORD, "submitted").status_code == 200
    r = client.patch(f"{API}/assessments/{a1}/self", headers=auth(COORD), json={"self_rating": "partial"})
    assert r.status_code == 409                                   # مجمّد بعد الرفع
    assert t(COORD, "hod_approved").status_code == 409            # ليس من صلاحية المنسق
    assert t(HOD, "returned").status_code == 409                  # الإعادة بلا سبب
    assert t(HOD, "hod_approved").status_code == 200
    assert t(AUDITOR, "under_audit").status_code == 200
    assert client.patch(f"{API}/assessments/{a1}/external", headers=auth(COORD),
                        json={"ext_verdict": "conforms"}).status_code == 403
    r = client.patch(f"{API}/assessments/{a1}/external", headers=auth(AUDITOR),
                     json={"ext_verdict": "insufficient",
                           "ext_recommendation": "اعتماد الهيكل من مجلس الكلية"})
    assert r.status_code == 200
    assert t(AUDITOR, "audited").status_code == 200
    assert t(DEAN, "final").status_code == 200
    hist = client.get(f"{API}/assessments/{a1}", headers=auth(HOD)).json()["history"]
    assert [h["to_status"] for h in hist] == ["submitted", "hod_approved", "under_audit", "audited", "final"]


def test_reports(client, auth):
    deps = client.get(f"{API}/cycles/1/reports/departments", headers=auth(DEAN)).json()
    cs = next(d for d in deps if d["department_id"] == 1)
    assert float(cs["self_pct"]) == 100.0 and float(cs["ext_pct"]) == 0.0 and cs["needs_improvement_count"] == 1
    assert [d["department_id"] for d in client.get(f"{API}/cycles/1/reports/departments",
                                                    headers=auth(COORD)).json()] == [1]
    gaps = client.get(f"{API}/cycles/1/reports/gaps", headers=auth(DEAN)).json()
    assert gaps[0]["assessment_id"] == STATE["a1"]
    assert gaps[0]["ext_verdict"] == "insufficient" and gaps[0]["eff_rating"] == "partial"
    rank = client.get(f"{API}/cycles/1/reports/ranking", headers=auth(DEAN)).json()
    assert rank["best"] and rank["worst"]
    level = client.get(f"{API}/cycles/1/reports/pass-level", headers=auth(DEAN)).json()
    assert level["passed"] is None and level["threshold_pct"] == 89.5     # التقييم الخارجي لم يكتمل
    plans = client.get(f"{API}/cycles/1/reports/final-plans", headers=auth(DEAN)).json()
    r1 = next(p for p in plans if p["assessment_id"] == STATE["a1"])        # جزئي بحكم «غير كافٍ» → يتطلب خطة
    assert r1["partially_implemented"] and r1["evidence_issue"] and not r1["not_implemented"]
    assert all(p["department_id"] == 1 for p in client.get(f"{API}/cycles/1/reports/final-plans", headers=auth(COORD)).json())


def test_na_requires_star_rule(client, auth):
    """«لا ينطبق» مرفوض لمتطلب بلا شرط نجمة (قاعدة النموذج)."""
    cs = client.get(f"{API}/cycles/1/assessments", headers=auth(COORD)).json()
    r2 = next(a for a in cs if a["requirement_code"] == "R2")
    assert r2["na_ok"] is False
    r = client.patch(f"{API}/assessments/{r2['id']}/self", headers=auth(COORD), json={"self_rating": "na"})
    assert r.status_code == 409 and "لا ينطبق" in r.json()["detail"]


def test_improvement_plan_module(client, auth):
    """خطة تحسين: تتولد تلقائياً، يعبّئها القسم، يحكم المقيّم، تُنفَّذ، ويُتحقق منها طرف مستقل."""
    from datetime import date, timedelta
    a1 = STATE["a1"]
    plans = client.get(f"{API}/cycles/1/actions", headers=auth(COORD)).json()
    plan = next(p for p in plans if p["assessment_id"] == a1)       # تولدت عند الاعتماد النهائي (جزئي)
    assert plan["source"] == "auto" and plan["status"] == "draft" and plan["requirement_code"] == "R1"
    pid = plan["id"]
    t = lambda who, to, **kw: client.post(f"{API}/actions/{pid}/transition", headers=auth(who),
                                          json={"to_status": to, **kw})
    patch = lambda who, **kw: client.patch(f"{API}/actions/{pid}", headers=auth(who), json=kw)

    r = t(COORD, "submitted")
    assert r.status_code == 409 and "الإجراءات التحسينية" in r.json()["detail"]
    past = (date.today() - timedelta(days=10)).isoformat()
    assert patch(COORD, action_ar="اعتماد الهيكل من مجلس الكلية وتعميمه", owner_id=1, due_date=past).status_code == 200
    assert t(COORD, "submitted").status_code == 200
    assert patch(COORD, action_ar="تعديل بعد الرفع").status_code == 409
    r = t("auditor.cs@ynb", "approved")
    assert r.status_code == 409 and "الفصل بين المهام" in r.json()["detail"]
    assert t(AUDITOR, "revision", comment="حدد آلية التعميم").status_code == 409        # بلا حكم ملاءمة
    assert t(AUDITOR, "revision", comment="حدد آلية التعميم", adequacy="insufficient").status_code == 200
    assert patch(COORD, action_ar="اعتماد الهيكل وتعميمه بخطاب رسمي على المنسوبين").status_code == 200
    assert t(COORD, "submitted").status_code == 200
    assert t(AUDITOR, "approved").status_code == 200

    overdue = client.get(f"{API}/cycles/1/actions?overdue=true", headers=auth(DEAN)).json()
    assert any(p["id"] == pid for p in overdue)
    assert patch(COORD, due_date=(date.today() + timedelta(days=30)).isoformat()).status_code == 409
    assert patch(HOD, due_date=(date.today() + timedelta(days=30)).isoformat()).status_code == 200

    r = t(COORD, "completed")
    assert r.status_code == 409 and "شاهد التحسين" in r.json()["detail"]
    assert patch(COORD, completion_evidence_id=STATE["ev"], progress_pct=80).status_code == 200
    assert t(COORD, "completed").status_code == 200
    assert t(COORD, "verified").status_code == 409                                       # ليس من دوره
    assert t(AUDITOR, "approved").status_code == 409                                     # رفض بلا سبب
    assert t(AUDITOR, "approved", comment="الخطاب غير موقع").status_code == 200
    assert t(COORD, "completed").status_code == 200
    assert t(DEAN, "verified", comment="تم التحقق من الخطاب الموقع").status_code == 200

    d = client.get(f"{API}/actions/{pid}", headers=auth(HOD)).json()
    assert d["status"] == "verified" and d["adequacy"] == "suitable" and d["verified_by_name"] == "وكيل الجودة"
    assert [h["to_status"] for h in d["history"]] == ["submitted", "revision", "submitted", "approved",
                                                     "completed", "approved", "completed", "verified"]
    summary = client.get(f"{API}/cycles/1/reports/improvement", headers=auth(DEAN)).json()
    cs = next(r for r in summary if r["department_id"] == 1)
    assert cs["verified"] == 1 and float(cs["closure_pct"]) == 100.0
    assert client.post(f"{API}/cycles/1/actions/generate", headers=auth(DEAN)).json()["generated"] == 0
    assert client.get(f"{API}/actions/{pid}", headers=auth(COORD_EE)).status_code == 404


def test_export_endpoint_guards(client, auth):
    assert client.get(f"{API}/cycles/1/export/official", headers=auth(COORD)).status_code == 403
    r = client.get(f"{API}/cycles/1/export/official", headers=auth(DEAN))
    assert r.status_code == 200 and r.content[:2] == b"PK"      # بلا قالب رسمي: ملف Excel بديل
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["ملخص الأقسام", "الاستيفاء حسب المعيار", "التقديرات التفصيلية", "خطط التحسين"]


def test_close_cycle_guarded(client, auth):
    r = client.post(f"{API}/cycles/1/close", headers=auth(DEAN))
    assert r.status_code == 409 and "لم تصل" in r.json()["detail"]
    assert client.post(f"{API}/cycles/1/close", headers=auth(HOD)).status_code == 409


def test_frontend_support_endpoints(client, auth):
    deps = client.get(f"{API}/departments", headers=auth(COORD)).json()
    assert [d["code"] for d in deps] == ["CS", "EE"]
    me = client.get(f"{API}/me", headers=auth(COORD)).json()
    assert me["institution_name"] and me["grants"][0]["role"] == "dept_coordinator"
    stds = client.get(f"{API}/cycles/1/reports/standards", headers=auth(DEAN)).json()
    assert stds and all("standard_title" in s for s in stds)
    assert client.get("/app/").status_code == 200


def test_requirement_scope_exclusion(client, auth):
    """متطلب خارج اختصاص القسم: يقرره وكيل الجودة بمسوّغ، فيختفي من تقديرات القسم ومؤشراته، ويعود عند الإلغاء."""
    tree = client.get(f"{API}/departments/2/applicability", headers=auth(COORD_EE)).json()
    assert tree["can_edit"] is False and tree["excluded_count"] == 0
    r3 = next(r for r in tree["requirements"] if r["requirement_code"] == "R3")
    assert client.get(f"{API}/departments/2/applicability", headers=auth(COORD)).status_code == 404
    body = {"requirement_ids": [r3["requirement_id"]], "reason_ar": "ليس من اختصاص القسم"}
    url = f"{API}/departments/2/applicability/exclusions"
    assert client.post(url, headers=auth(COORD_EE), json=body).status_code == 403
    r = client.post(url, headers=auth(DEAN), json=body)                    # خارج شرط النجمة → يلزم تأكيد
    assert r.status_code == 403 and "R3" in r.json()["detail"]
    r = client.post(url, headers=auth(DEAN), json={**body, "outside_model_rule": True})
    assert r.status_code == 201 and r.json()["excluded"] == 1

    ee = client.get(f"{API}/cycles/1/assessments", headers=auth(COORD_EE)).json()
    assert "R3" not in {a["requirement_code"] for a in ee} and len(ee) == 2
    subs = client.get(f"{API}/cycles/1/reports/sub-standards?department_id=2", headers=auth(DEAN)).json()
    assert "S1.2" not in {s["sub_standard_code"] for s in subs}            # المعيار الفرعي كله خارج النطاق
    tree = client.get(f"{API}/departments/2/applicability", headers=auth(DEAN)).json()
    x = next(r for r in tree["requirements"] if r["requirement_code"] == "R3")
    assert tree["excluded_count"] == 1 and x["within_model_rule"] is False and x["reason_ar"] == "ليس من اختصاص القسم"

    assert client.delete(f"{url}/{r3['requirement_id']}", headers=auth(DEAN)).status_code == 200
    ee = client.get(f"{API}/cycles/1/assessments", headers=auth(COORD_EE)).json()
    assert "R3" in {a["requirement_code"] for a in ee}
    assert client.delete(f"{url}/{r3['requirement_id']}", headers=auth(DEAN)).status_code == 404


def test_calendar(client, auth):
    """تقويم الإنجاز: يرفعه وكيل الجودة (إدخال أو ملف)، ويحتسب إنجاز الأقسام، والموعد الأخص يغلب."""
    url = f"{API}/cycles/1/calendar"
    body = {"target_code": "S1.1", "due_on": "2026-11-30", "starts_on": "2026-10-01", "note_ar": "قبل نهاية الفصل"}
    assert client.post(url, headers=auth(COORD), json=body).status_code == 403
    assert client.post(url, headers=auth(DEAN), json={**body, "target_code": "X9"}).status_code == 404
    assert client.post(url, headers=auth(DEAN), json=body).status_code == 201
    assert client.post(url, headers=auth(DEAN), json={**body, "due_on": "2026-12-15"}).status_code == 201   # تحديث لا تكرار
    cal = client.get(url, headers=auth(DEAN)).json()
    s11 = next(d for d in cal["deadlines"] if d["target_code"] == "S1.1")
    assert cal["can_edit"] and len(cal["deadlines"]) == 1 and s11["due_on"] == "2026-12-15"
    assert s11["total"] == 4 and s11["target_kind"] == "sub_standard"          # R1، R2 × قسمين
    mine = client.get(url, headers=auth(COORD)).json()["deadlines"][0]
    assert mine["total"] == 2 and mine["done"] >= 1                           # R1 للحاسب معتمد نهائياً

    csv_text = "﻿الرمز,الاسم,البدء,الانتهاء,القسم,ملاحظات\nR3,,,20/01/2027,EE,\nS1.2,,,2027-02-01,,\nZZ,,,2027-01-01,,\nR1,,,غداً,,\n"
    r = client.post(f"{url}/import", headers=auth(DEAN), files={"file": ("cal.csv", csv_text.encode(), "text/csv")})
    out = r.json()
    assert r.status_code == 200 and out["imported"] == 2 and len(out["errors"]) == 2
    a_ee = client.get(f"{API}/cycles/1/assessments", headers=auth(COORD_EE)).json()
    assert next(a for a in a_ee if a["requirement_code"] == "R3")["deadline"] == "2027-01-20"   # المتطلب + القسم يغلب
    a_cs = client.get(f"{API}/cycles/1/assessments", headers=auth(COORD)).json()
    assert next(a for a in a_cs if a["requirement_code"] == "R3")["deadline"] == "2027-02-01"
    assert client.get(f"{url}/template", headers=auth(DEAN)).text.lstrip("﻿").startswith("الرمز")
    did = s11["id"]
    assert client.delete(f"{API}/calendar/{did}", headers=auth(COORD)).status_code == 403
    assert client.delete(f"{API}/calendar/{did}", headers=auth(DEAN)).status_code == 200


def test_evidence_links_policy_and_folders(client, auth):
    """الشواهد بالروابط: سياسة المنشأة، النطاقات المسموحة، مجلدات الأقسام."""
    s = client.get(f"{API}/settings/evidence", headers=auth(COORD)).json()
    assert s["attachment_mode"] == "both" and s["can_edit"] is False
    assert client.put(f"{API}/settings/evidence", headers=auth(COORD), json={"attachment_mode": "links"}).status_code == 403
    r = client.put(f"{API}/settings/evidence", headers=auth(DEAN),
                   json={"attachment_mode": "links", "allowed_link_domains": ["https://tvtc-my.sharepoint.com/"]})
    assert r.status_code == 200 and r.json()["allowed_link_domains"] == ["tvtc-my.sharepoint.com"]

    ok = client.post(f"{API}/evidence", headers=auth(COORD), json={"title_ar": "خطة الإرشاد", "kind": "link", "department_id": 1,
                     "url": "https://tvtc-my.sharepoint.com/personal/q/Documents/plan.pdf"})
    assert ok.status_code == 201
    bad = client.post(f"{API}/evidence", headers=auth(COORD), json={"title_ar": "رابط شخصي", "kind": "link", "department_id": 1,
                      "url": "https://drive.google.com/file/d/x"})
    assert bad.status_code == 409 and "النطاقات المعتمدة" in bad.json()["detail"]
    f = client.post(f"{API}/evidence", headers=auth(COORD), json={"title_ar": "ملف مرفوع", "kind": "file", "department_id": 1}).json()
    up = client.post(f"{API}/evidence/{f['id']}/versions", headers=auth(COORD), files={"file": ("a.pdf", b"%PDF-1.4", "application/pdf")})
    assert up.status_code == 403 and "الروابط فقط" in up.json()["detail"]

    url = f"{API}/cycles/1/folders"
    assert client.put(url, headers=auth(COORD), json={"department_id": 1, "url": "https://tvtc-my.sharepoint.com/f"}).status_code == 403
    assert client.put(url, headers=auth(DEAN), json={"department_id": 1, "url": "https://evil.example.com/f"}).status_code == 403
    assert client.put(url, headers=auth(DEAN), json={"department_id": 1, "url": "https://tvtc-my.sharepoint.com/root"}).status_code == 200
    assert client.put(url, headers=auth(DEAN), json={"department_id": 1, "target_code": "S1.1",
                                                     "url": "https://tvtc-my.sharepoint.com/s11"}).status_code == 200
    fs = client.get(f"{url}?department_id=1", headers=auth(COORD)).json()
    assert {(x["sub_standard_code"], x["url"].rsplit("/", 1)[1]) for x in fs} == {(None, "root"), ("S1.1", "s11")}
    assert client.get(f"{url}?department_id=1", headers=auth(COORD_EE)).json() == []
    assert client.get(f"{API}/evidence/links/problems", headers=auth(COORD)).status_code == 200
    assert client.put(f"{API}/settings/evidence", headers=auth(DEAN),
                      json={"attachment_mode": "both", "allowed_link_domains": []}).status_code == 200


def test_users_and_guides_admin(client, auth):
    """إدارة المستخدمين والأدلة: لوكيل الجودة فقط، مع حماية دور مدير النظام والحساب الذاتي."""
    assert client.get(f"{API}/users", headers=auth(COORD)).status_code == 403
    data = client.get(f"{API}/users", headers=auth(DEAN)).json()
    assert data["can_grant_admin"] is False and any(u["email"] == "dean@ynb" for u in data["users"])
    r = client.post(f"{API}/users", headers=auth(DEAN), json={"email": "New.Trainer@ynb", "full_name_ar": "مدرب جديد"})
    assert r.status_code == 201
    uid = r.json()["id"]
    assert client.post(f"{API}/users", headers=auth(DEAN), json={"email": "new.trainer@ynb", "full_name_ar": "مكرر"}).status_code == 409
    assert client.post(f"{API}/users/{uid}/roles", headers=auth(DEAN), json={"role": "hod"}).status_code == 403          # بلا قسم
    assert client.post(f"{API}/users/{uid}/roles", headers=auth(DEAN), json={"role": "system_admin"}).status_code == 403
    rr = client.post(f"{API}/users/{uid}/roles", headers=auth(DEAN), json={"role": "faculty", "department_id": 1})
    assert rr.status_code == 201
    me = client.get(f"{API}/me", headers=auth("new.trainer@ynb")).json()
    assert me["grants"][0]["role"] == "faculty"
    dean_id = next(u["id"] for u in data["users"] if u["email"] == "dean@ynb")
    assert client.patch(f"{API}/users/{dean_id}", headers=auth(DEAN), json={"is_active": False}).status_code == 403
    assert client.patch(f"{API}/users/{uid}", headers=auth(DEAN), json={"is_active": False}).status_code == 200
    assert client.delete(f"{API}/roles/{rr.json()['id']}", headers=auth(DEAN)).status_code == 200

    assert client.post(f"{API}/guides", headers=auth(COORD), json={"title_ar": "دليل", "url": "https://x.sa/a.pdf"}).status_code == 403
    g = client.post(f"{API}/guides", headers=auth(DEAN), json={"title_ar": "دليل رئيس القسم", "url": "https://tvtc.gov.sa/guide.pdf"})
    assert g.status_code == 201
    lst = client.get(f"{API}/guides", headers=auth(COORD)).json()
    assert lst["can_edit"] is False and lst["guides"][0]["title_ar"] == "دليل رئيس القسم"
    assert client.delete(f"{API}/guides/{g.json()['id']}", headers=auth(DEAN)).status_code == 200


def test_quality_works(client, auth):
    """أعمال الجودة: جدول مهام أسبوعي مرتبط بالمتطلبات؛ تسليم القسم ينعكس شاهداً على تقديراته دون تكرار."""
    base = f"{API}/cycles/1/quality-tasks"
    assert client.post(f"{base}/seed", headers=auth(COORD), json={"semester_start": "2026-08-30"}).status_code == 403
    r = client.post(f"{base}/seed", headers=auth(DEAN), json={"semester_start": "2026-08-31"})   # الاثنين → يُحسب من الأحد
    assert r.status_code == 201 and r.json()["created"] == 26
    assert client.post(f"{base}/seed", headers=auth(DEAN), json={"semester_start": "2026-08-30"}).status_code == 403
    tasks = client.get(base, headers=auth(DEAN)).json()["tasks"]
    first = tasks[0]
    assert first["week_no"] == 3 and first["due_on"] == "2026-09-17" and "الثالث" in first["due_hint_ar"]

    t = client.post(base, headers=auth(DEAN), json={"week_no": 5, "title_ar": "تقرير تجريبي للربط"}).json()["id"]
    assert client.put(f"{API}/quality-tasks/{t}/links", headers=auth(DEAN), json={"requirement_codes": ["ZZ9"]}).status_code == 404
    assert client.put(f"{API}/quality-tasks/{t}/links", headers=auth(COORD), json={"requirement_codes": ["R3"]}).status_code == 403
    assert client.put(f"{API}/quality-tasks/{t}/links", headers=auth(DEAN),
                      json={"sub_standard_codes": ["S1.1"], "requirement_codes": ["R3"]}).status_code == 200

    sub = {"department_id": 1, "url": "https://tvtc-my.sharepoint.com/personal/q/Documents/exam.pdf", "note_ar": "تقرير الأسبوع"}
    assert client.post(f"{API}/quality-tasks/{t}/submit", headers=auth(COORD_EE), json=sub).status_code == 403
    out = client.post(f"{API}/quality-tasks/{t}/submit", headers=auth(COORD), json=sub).json()
    r3 = [x for x in client.get(f"{API}/cycles/1/assessments", headers=auth(COORD)).json() if x["requirement_code"] == "R3"]
    assert out["status"] == "submitted" and out["linked"] + out["skipped"] == len(r3)
    if out["linked"]:
        assert r3[0]["evidence_count"] >= 1                 # انعكس الشاهد على تقدير المتطلب
    mine = next(x for x in client.get(base, headers=auth(COORD)).json()["tasks"] if x["id"] == t)
    assert mine["submissions"]["1"]["status"] == "submitted" and {l["code"] for l in mine["links"]} == {"S1.1", "R3"}
    assert client.post(f"{API}/quality-tasks/{t}/review", headers=auth(DEAN),
                       json={"department_id": 1, "decision": "returned"}).status_code == 403   # الإعادة تتطلب سبباً
    assert client.post(f"{API}/quality-tasks/{t}/review", headers=auth(DEAN),
                       json={"department_id": 1, "decision": "approved"}).json()["status"] == "approved"
    na = client.post(f"{API}/quality-tasks/{first['id']}/submit", headers=auth(COORD), json={"department_id": 1, "not_applicable": True})
    assert na.status_code == 403
    assert client.post(f"{API}/quality-tasks/{first['id']}/submit", headers=auth(COORD),
                       json={"department_id": 1, "not_applicable": True, "note_ar": "القسم لا يطبق التدريب عن بعد"}).json()["status"] == "na"
    assert len(client.get(f"{base}/catalog", headers=auth(COORD)).json()) >= 1
    assert client.delete(f"{API}/quality-tasks/{t}", headers=auth(COORD)).status_code == 403
